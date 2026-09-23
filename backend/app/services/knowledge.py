"""Shop knowledge: the guided sections, file-to-text extraction, and the prompt text."""

import io
from dataclasses import dataclass

from app.models import KnowledgeEntry, KnowledgeSection
from app.services.llm import LLMClient, LLMError


@dataclass(frozen=True)
class SectionInfo:
    key: KnowledgeSection
    label: str
    hint: str
    placeholder: str
    # Single sections are one text box in the panel; multiple ones are a list of entries.
    multiple: bool


# Order here is the order the panel shows them and the order the bot reads them.
SECTIONS: list[SectionInfo] = [
    SectionInfo(
        KnowledgeSection.ABOUT,
        "معرفی فروشگاه",
        "فروشگاه شما چی می‌فروشه و چه چیزی باید درباره‌اش بدونیم؟",
        "مثال: مزون مهسا، تولیدکننده‌ی مانتو و شومیز زنانه از سال ۱۳۹۸. همه‌ی لباس‌ها دوخت خودمونه.",
        False,
    ),
    SectionInfo(
        KnowledgeSection.SHIPPING,
        "ارسال",
        "روش‌های ارسال، هزینه، زمان تحویل، ارسال به شهرستان و خارج از کشور",
        "مثال: ارسال با پست پیشتاز ۳ تا ۵ روز کاری، هزینه ۸۵ هزار تومان. خرید بالای ۲ میلیون رایگان.",
        False,
    ),
    SectionInfo(
        KnowledgeSection.PAYMENT,
        "پرداخت",
        "روش‌های پرداخت و نحوه‌ی ثبت سفارش",
        "مثال: کارت‌به‌کارت یا درگاه سایت. پرداخت در محل فقط برای تهران.",
        False,
    ),
    SectionInfo(
        KnowledgeSection.RETURNS,
        "مرجوعی و تعویض",
        "شرایط و مهلت تعویض یا مرجوعی",
        "مثال: تا ۷ روز بعد از تحویل، در صورت سالم بودن و داشتن اتیکت، امکان تعویض سایز هست.",
        False,
    ),
    SectionInfo(
        KnowledgeSection.CONTACT,
        "راه‌های ارتباط و ساعت کاری",
        "سایت، تلفن، آدرس حضوری، ساعت پاسخگویی",
        "مثال: سایت mahsa.shop، پشتیبانی ۱۰ صبح تا ۸ شب، جمعه‌ها تعطیل.",
        False,
    ),
    SectionInfo(
        KnowledgeSection.FAQ,
        "سوال‌های پرتکرار",
        "سوال‌هایی که مشتری‌ها زیاد می‌پرسن و جوابشون",
        "سوال: جنس پارچه چیه؟ / جواب: کتان ترک درجه یک",
        True,
    ),
    SectionInfo(
        KnowledgeSection.FREE_TEXT,
        "توضیحات دیگر",
        "هر اطلاعات دیگری که ربات باید بدونه",
        "مثال: تخفیف ۱۵٪ مخصوص خرید اول با کد FIRST15",
        True,
    ),
    SectionInfo(
        KnowledgeSection.FILE,
        "فایل‌های آپلودشده",
        "متن استخراج‌شده از فایل‌ها. قابل ویرایش است.",
        "",
        True,
    ),
]
SECTION_BY_KEY = {s.key: s for s in SECTIONS}

TEXT_TYPES = {".txt", ".md", ".csv"}
IMAGE_TYPES = {
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
    ".png": "image/png",
    ".webp": "image/webp",
}
SUPPORTED_EXTENSIONS = sorted(TEXT_TYPES | {".pdf", ".docx"} | IMAGE_TYPES.keys())


class ExtractionError(Exception):
    """User-facing (Persian) explanation of why a file could not be read."""


async def extract_text(filename: str, data: bytes, llm: LLMClient) -> str:
    ext = "." + filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if ext in TEXT_TYPES:
        text = _decode(data)
    elif ext == ".pdf":
        text = _pdf_text(data)
    elif ext == ".docx":
        text = _docx_text(data)
    elif ext in IMAGE_TYPES:
        if not llm.configured:
            raise ExtractionError("خواندن عکس فعلا فعال نیست. لطفا متن را تایپ کنید.")
        try:
            text = await llm.extract_text_from_image(data, IMAGE_TYPES[ext])
        except LLMError as exc:
            raise ExtractionError("خواندن عکس با خطا مواجه شد. دوباره تلاش کنید.") from exc
    else:
        allowed = "، ".join(SUPPORTED_EXTENSIONS)
        raise ExtractionError(f"این نوع فایل پشتیبانی نمی‌شود. فرمت‌های مجاز: {allowed}")

    text = "\n".join(line.rstrip() for line in text.splitlines()).strip()
    if not text:
        raise ExtractionError(
            "متنی در این فایل پیدا نشد. اگر PDF اسکن‌شده است، عکس صفحه‌ها را آپلود کنید."
        )
    return text


def _decode(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-16", "cp1256"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ExtractionError("کدگذاری فایل متنی قابل تشخیص نیست. فایل را با UTF-8 ذخیره کنید.")


def _pdf_text(data: bytes) -> str:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(io.BytesIO(data))
        return "\n\n".join(page.extract_text() or "" for page in reader.pages)
    except (PdfReadError, ValueError) as exc:
        raise ExtractionError("فایل PDF خراب است یا رمز دارد.") from exc


def _docx_text(data: bytes) -> str:
    import docx

    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:  # python-docx raises several unrelated types for bad files
        raise ExtractionError("فایل Word قابل خواندن نیست.") from exc
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(parts)


def total_chars(entries: list[KnowledgeEntry]) -> int:
    return sum(len(e.title) + len(e.content) for e in entries)


def render_for_prompt(entries: list[KnowledgeEntry]) -> str:
    """Group entries under section headings, in the canonical section order."""
    blocks: list[str] = []
    for section in SECTIONS:
        items = [e for e in entries if e.section == section.key and e.content.strip()]
        if not items:
            continue
        lines = [f"## {section.label}"]
        for item in items:
            if item.title:
                lines.append(f"### {item.title}")
            lines.append(item.content.strip())
        blocks.append("\n".join(lines))
    return "\n\n".join(blocks)
